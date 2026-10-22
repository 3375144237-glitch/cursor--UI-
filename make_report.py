#!/usr/bin/env python3
"""根据走查记录生成：标注截图、问题对比图、Excel 问题表（沿用团队模板列）。

用法：
  .venv/bin/python make_report.py reviews/maiqi/issues.yaml
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import date
from io import BytesIO
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image

from src.annotate import annotate, color, crop_pair

ROOT = Path(__file__).resolve().parent
HEADERS = ["设备", "截图", "问题描述", "优先级", "修复状态", "备注"]
WIDTHS = [14, 40, 60, 11, 13, 30]          # 列宽（字符）
HEADER_FILL = "A9B4F7"
PRIORITIES = ["高", "中", "低"]
STATUSES = ["待修复", "修复中", "已修复", "不修复", "待确认"]
THUMB_W = 260                              # 截图列中缩略图宽度（像素）
CIRCLED = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳"
PRIORITY_FILL = {"高": "FFD6D6", "中": "FFE9C7", "低": "E6F4EA"}


def png(img: Image.Image) -> BytesIO:
    buf = BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf


def design_image(cfg_dir: Path, node: str | None) -> Image.Image | None:
    if not node:
        return None
    index = json.loads((ROOT / "screenshots/design/maiqi/index.json").read_text(encoding="utf-8"))
    entry = index.get(node)
    return Image.open(ROOT / "screenshots/design/maiqi" / entry["file"]) if entry else None


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("issues", help="走查记录 yaml")
    ap.add_argument("--out", default=None, help="输出目录，默认 output/<项目>_<日期>")
    args = ap.parse_args()

    src = Path(args.issues)
    data = yaml.safe_load(src.read_text(encoding="utf-8"))
    shots_dir = Path(data["screenshot_dir"]).expanduser()
    out = Path(args.out) if args.out else ROOT / "output" / f"{data['project']}_{date.today():%Y%m%d}"
    (out / "标注截图").mkdir(parents=True, exist_ok=True)
    (out / "问题对比").mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "UI走查问题"
    thin = Side(style="thin", color="D0D4E4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(HEADERS)
    ws.row_dimensions[1].height = 34
    for cell, width in zip(ws[1], WIDTHS):
        cell.font = Font(bold=True, color="FFFFFF", size=14)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        ws.column_dimensions[cell.column_letter].width = width
    ws["D1"].comment = Comment("可选：" + " / ".join(PRIORITIES), "UI走查")
    ws["E1"].comment = Comment("可选：" + " / ".join(STATUSES), "UI走查")

    dv_p = DataValidation(type="list", formula1='"' + ",".join(PRIORITIES) + '"', allow_blank=True)
    dv_s = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv_p)
    ws.add_data_validation(dv_s)

    detail = wb.create_sheet("问题对比图")
    detail.append(["页面", "编号", "设计稿 vs 实机（局部）", "问题描述"])
    for cell, width in zip(detail[1], [20, 8, 90, 50]):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        detail.column_dimensions[cell.column_letter].width = width

    row = 2
    drow = 2
    total = 0
    for page in data["pages"]:
        issues = page.get("issues") or []
        if not issues:
            continue
        shot = Image.open(shots_dir / page["shot"])
        design = design_image(src.parent, page.get("node"))
        page_offset = page.get("design_offset", [0, -12])

        marked = annotate(shot, [i["box"] for i in issues])
        marked_path = out / "标注截图" / f"{page['page']}.png"
        marked.save(marked_path)

        thumb_h = round(marked.height * THUMB_W / marked.width)
        thumb = marked.resize((THUMB_W, thumb_h), Image.LANCZOS)
        n = len(issues)
        # 合并单元格的总高度至少容纳缩略图；每行再保证文字放得下
        per_row_pt = max(thumb_h * 0.75 / n + 4, 48)

        first = row
        for k, issue in enumerate(issues, 1):
            ws.cell(row=row, column=3, value=f"{CIRCLED[k - 1]} {issue['desc']}")
            ws.cell(row=row, column=4, value=issue.get("priority", "中"))
            ws.cell(row=row, column=5, value=issue.get("status", "待修复"))
            ws.cell(row=row, column=6, value=issue.get("note", ""))
            ws.row_dimensions[row].height = per_row_pt
            for col in range(1, 7):
                c = ws.cell(row=row, column=col)
                c.border = border
                c.alignment = Alignment(vertical="center", wrap_text=True,
                                        horizontal="center" if col in (1, 4, 5) else "left")
            ws.cell(row=row, column=3).font = Font(size=11)
            pr = issue.get("priority", "中")
            if pr in PRIORITY_FILL:
                ws.cell(row=row, column=4).fill = PatternFill("solid", fgColor=PRIORITY_FILL[pr])
            dv_p.add(ws.cell(row=row, column=4))
            dv_s.add(ws.cell(row=row, column=5))

            if design is not None:
                off = issue.get("design_offset", page_offset)
                pair = crop_pair(design, shot, issue["box"], k, design_offset=tuple(off))
                pair_path = out / "问题对比" / f"{page['page']}_{k}.png"
                pair.save(pair_path)
                scale = min(1.0, 620 / pair.width)
                small = pair.resize((round(pair.width * scale), round(pair.height * scale)), Image.LANCZOS)
                detail.cell(row=drow, column=1, value=page["page"])
                detail.cell(row=drow, column=2, value=CIRCLED[k - 1])
                detail.cell(row=drow, column=4, value=issue["desc"])
                for col in (1, 2, 4):
                    detail.cell(row=drow, column=col).alignment = Alignment(vertical="center", wrap_text=True)
                xl = XLImage(png(small))
                xl.width, xl.height = small.width, small.height
                detail.add_image(xl, f"C{drow}")
                detail.row_dimensions[drow].height = small.height * 0.75 + 8
                drow += 1
            row += 1
            total += 1

        last = row - 1
        ws.cell(row=first, column=1, value=f"{data['device']}\n{page['page']}")
        if last > first:
            ws.merge_cells(start_row=first, start_column=1, end_row=last, end_column=1)
            ws.merge_cells(start_row=first, start_column=2, end_row=last, end_column=2)
        ws.cell(row=first, column=1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        xl = XLImage(png(thumb))
        xl.width, xl.height = thumb.width, thumb.height
        ws.add_image(xl, f"B{first}")

    ws.freeze_panes = "A2"
    xlsx = out / f"{data['project']}_UI走查问题表_{date.today():%Y%m%d}.xlsx"
    wb.save(xlsx)
    print(f"问题数 {total}  页面数 {sum(1 for p in data['pages'] if p.get('issues'))}")
    print(f"问题表   {xlsx}")
    print(f"标注截图 {out / '标注截图'}")
    print(f"问题对比 {out / '问题对比'}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
