"""生成带对比小图的 Excel 问题表。"""
from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path

import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill

from .compare import Region, crop_pair

HEADERS = ["序号", "页面", "问题位置", "设计稿 vs 实机", "问题描述", "严重程度", "确认（是/否）", "备注"]
COL_WIDTHS = [6, 14, 18, 56, 42, 10, 13, 22]
CROP_WIDTH = 380


def _png_bytes(image: np.ndarray) -> BytesIO:
    ok, buf = cv2.imencode(".png", image)
    if not ok:
        raise RuntimeError("图片编码失败")
    return BytesIO(buf.tobytes())


def build_report(
    page: str,
    design: np.ndarray,
    actual: np.ndarray,
    regions: list[Region],
    output: str | Path,
    descriptions: dict[int, str] | None = None,
) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "走查问题"

    ws.append(HEADERS)
    for cell, width in zip(ws[1], COL_WIDTHS):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="3D5A80")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = width

    for idx, region in enumerate(regions, 1):
        pair = crop_pair(design, actual, region)
        scale = CROP_WIDTH / pair.shape[1]
        resized = cv2.resize(pair, (CROP_WIDTH, int(pair.shape[0] * scale)), interpolation=cv2.INTER_AREA)

        row = idx + 1
        severity = "高" if region.peak_diff >= 80 and region.area > 3000 else "中" if region.peak_diff >= 50 else "低"
        desc = (descriptions or {}).get(idx) or (
            f"区域 ({region.x},{region.y}) 尺寸 {region.w}x{region.h}，"
            f"{'实机缺失该元素' if region.kind == 'missing' else '视觉样式不一致'}，"
            f"峰值色差 {region.peak_diff}/255"
        )
        ws.cell(row=row, column=1, value=idx)
        ws.cell(row=row, column=2, value=page)
        ws.cell(row=row, column=3, value=f"x={region.x}, y={region.y}")
        ws.cell(row=row, column=5, value=desc)
        ws.cell(row=row, column=6, value=severity)
        ws.cell(row=row, column=7, value="")
        ws.cell(row=row, column=8, value="")
        for col in range(1, 9):
            ws.cell(row=row, column=col).alignment = Alignment(vertical="center", wrap_text=True)

        img = XLImage(_png_bytes(resized))
        img.width, img.height = CROP_WIDTH, resized.shape[0]
        ws.add_image(img, f"D{row}")
        ws.row_dimensions[row].height = resized.shape[0] * 0.75 + 10

    ws.freeze_panes = "A2"
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output


def save_pair_image(design: np.ndarray, actual: np.ndarray, path: str | Path) -> Path:
    """输出左右并排的对比大图（左侧设计稿，右侧实机）。"""
    h = min(design.shape[0], actual.shape[0])
    gap = np.full((h, 8, 3), 200, np.uint8)
    merged = np.hstack([design[:h], gap, actual[:h]])
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    cv2.imwrite(str(path), merged)
    return path


def stamp() -> str:
    return date.today().strftime("%Y%m%d")
